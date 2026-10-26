#!/usr/bin/env python3
"""Generate an Excel evidence workbook for one case (and a global index).

Usage: python tools/build_evidence_xlsx.py [CASE_ID]
Default case: AZ-100-05-003. Output: evidence/<CASE_ID>.xlsx + evidence/index.xlsx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "evidence"
OUT.mkdir(exist_ok=True)
CASE = (sys.argv[1] if len(sys.argv) > 1 else "AZ-100-05-003").strip()

HEAD_FILL = PatternFill("solid", fgColor="1F4E5A")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="D5DDD8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {
    "exact": PatternFill("solid", fgColor="E4F4EA"),
    "blockEnd": PatternFill("solid", fgColor="FFF4DC"),
    "manual": PatternFill("solid", fgColor="F1EAFE"),
    "missing": PatternFill("solid", fgColor="FBE4E4"),
    "目視": PatternFill("solid", fgColor="F2F4F3"),
}
STATUS_LABEL = {"exact": "自動", "blockEnd": "近似", "manual": "人工確認", "missing": "未取得", "目視": "目視"}


def load_json(path: Path, default=None):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8-sig"))


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def build_case(case: str) -> Path:
    result = load_json(ROOT / "logs" / case / f"{case}-node-result.json", {}) or {}
    rows = (load_json(ROOT / "workflow" / f"workflow-{case}.json", {}) or {}).get("workflow", {}).get("procedure", {}).get("rawRows", [])
    evidence = {e["row"]: e for e in result.get("evidence", [])}
    summary = load_json(ROOT / "logs" / "evidence-summary.json", {}).get("cases", {}).get(case, {})
    evsum = summary.get("summary", {})

    wb = Workbook()
    ws = wb.active
    ws.title = "サマリ"
    ws["A1"] = "証跡サマリ"
    ws["A1"].font = Font(bold=True, size=16, color="1F4E5A")
    meta = [
        ("ケースID", case),
        ("実行状態", result.get("status", "—")),
        ("画面HC", evsum.get("hc", 0)),
        ("自動", evsum.get("exact", 0)),
        ("近似", evsum.get("blockEnd", 0)),
        ("人工確認", evsum.get("manual", 0)),
        ("未取得", evsum.get("missing", 0)),
        ("生成日時", result.get("ts", "")),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 36

    ws2 = wb.create_sheet("手順")
    headers = ["Excel行", "画面", "操作", "確認", "内容/期待", "証跡区分", "状態", "画像"]
    ws2.append(headers)
    style_header(ws2)
    widths = [8, 22, 28, 26, 40, 12, 12, 46]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for row in rows:
        op = " ".join(x for x in [row.get("operationType"), row.get("operationItem")] if x).strip()
        ver = " ".join(x for x in [row.get("verificationType"), row.get("verificationItem")] if x).strip()
        content = row.get("operationContent") or row.get("verificationContent") or ""
        if not (op or ver or content):
            continue
        ev = evidence.get(row.get("sourceRow"))
        status = "目視" if (row.get("evidenceType") or "") != "画面HC" else "missing"
        status = (ev or {}).get("status") or status
        label = STATUS_LABEL.get(status, status)
        ws2.cell(r, 1, row.get("sourceRow"))
        ws2.cell(r, 2, row.get("functionName") or "")
        ws2.cell(r, 3, op)
        ws2.cell(r, 4, ver)
        ws2.cell(r, 5, content)
        ws2.cell(r, 6, row.get("evidenceType") or "—")
        ws2.cell(r, 7, label)
        for c in range(1, 8):
            cell = ws2.cell(r, c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        if status in STATUS_FILL:
            ws2.cell(r, 7).fill = STATUS_FILL[status]
        img_name = (ev or {}).get("file")
        if img_name:
            img_path = ROOT / "logs" / case / img_name
            if img_path.exists():
                img = XLImage(str(img_path))
                img.width = 300
                img.height = 198
                ws2.add_image(img, f"H{r}")
                ws2.row_dimensions[r].height = 155
        if ws2.row_dimensions[r].height is None:
            ws2.row_dimensions[r].height = 34
        r += 1
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:H{r-1}"

    ws3 = wb.create_sheet("人工確認")
    ws3.append(["Excel行", "区分", "物件", "理由/内容", "確認"])
    style_header(ws3)
    for i, w in enumerate([8, 10, 40, 46, 10], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for i, m in enumerate((summary.get("manualSteps") or []), start=2):
        ws3.cell(i, 1, m.get("row"))
        ws3.cell(i, 2, m.get("phase"))
        ws3.cell(i, 3, m.get("item"))
        ws3.cell(i, 4, f"{m.get('reason')} {m.get('content') or ''}".strip())
        ws3.cell(i, 5, "")
        for c in range(1, 6):
            ws3.cell(i, c).alignment = Alignment(vertical="top", wrap_text=True)
            ws3.cell(i, c).border = BORDER
    ws3.freeze_panes = "A2"

    path = OUT / f"{case}.xlsx"
    wb.save(path)
    return path


def build_index():
    summary = load_json(ROOT / "logs" / "evidence-summary.json", {}) or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "ケース一覧"
    headers = ["ケースID", "実行", "画面HC", "自動", "近似", "人工確認", "未取得", "人工操作", "人工確認項目"]
    ws.append(headers)
    style_header(ws)
    for i, w in enumerate([18, 12, 10, 8, 8, 10, 8, 10, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, (cid, data) in enumerate(sorted(summary.get("cases", {}).items()), start=2):
        s = data.get("summary", {})
        ms = data.get("manualSteps", [])
        vals = [cid, data.get("status", ""), s.get("hc", 0), s.get("exact", 0), s.get("blockEnd", 0), s.get("manual", 0), s.get("missing", 0), sum(1 for m in ms if m.get("phase") == "操作"), sum(1 for m in ms if m.get("phase") == "確認")]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).border = BORDER
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(1, ws.max_row)}"
    path = OUT / "index.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    p = build_case(CASE)
    idx = build_index()
    print(f"wrote {p.relative_to(ROOT)}")
    print(f"wrote {idx.relative_to(ROOT)}")
